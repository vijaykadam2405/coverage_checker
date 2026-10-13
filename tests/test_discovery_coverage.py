"""Discovery reporting must measure evidence, not fabricate completeness."""
import json
import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

from crawler_coverage_auditor.analysis.dashboard import build_analysis
from crawler_coverage_auditor.analysis.inventory_segments import archive_year, archive_coverage
from crawler_coverage_auditor.reporting.html import write_html
from crawler_coverage_auditor.reporting.excel import write_excel
from crawler_coverage_auditor.discovery.pagination import pagination_series_key
from crawler_coverage_auditor.storage.database import AuditDatabase
from crawler_coverage_auditor.storage.models import DiscoveredURL
from scripts.refresh_reports import refresh
from crawler_coverage_auditor.auditor import CoverageAuditor
from crawler_coverage_auditor.config import AuditConfig
from crawler_coverage_auditor.discovery.browser import BrowserDiscovery
from crawler_coverage_auditor.discovery.static_links import Link


def row(path, match='EXACT_MATCH', validity='VALID_PAGE', **extra):
    url = 'https://example.test' + path
    return {'url': url, 'normalized_url': url, 'page_type': 'HTML', 'validity': validity,
            'match_status': match, 'discovery_methods': '["STATIC"]', 'min_depth': 1,
            **extra}


def area(analysis, name):
    return next(item for item in analysis['segments'] if item['Area'] == name)


def test_static_pagination_without_method_tag_and_pending_not_missing():
    analysis = build_analysis([row('/news?page=2'), row('/news?page=3', 'NOT_FOUND'),
                               row('/news?page=4', None, None), row('/news?page=5', 'API_ERROR')], {})
    pagination = area(analysis, 'Pagination')
    assert (pagination['Valid URLs'], pagination['Found'], pagination['Missing'], pagination['Unverified']) == (3, 1, 1, 1)
    assert pagination['Coverage %'] == 50
    detail = analysis['pagination'][0]
    assert detail['Candidates'] == 4 and detail['Awaiting website'] == 1
    assert detail['Website maximum'] is None and not detail['Maximum verified']
    assert detail['Missing ranges'] == '3'
    assert detail['Verification %'] == 66.67


def test_verified_boundary_is_preserved_and_invalid_probes_excluded():
    url = 'https://example.test/news?page=2'
    analysis = build_analysis([row('/news?page=2'), row('/news?page=3', 'NOT_FOUND'),
                               row('/news?page=4', None, 'HTTP_ERROR')], {'pagination': [
        {'template_url': url, 'parameter_name': 'page', 'series_key': pagination_series_key(url, 'page'),
         'status': 'VERIFIED', 'max_page': 3, 'observed_max': 3, 'evidence': 'Fixture boundary'}]})
    item = analysis['pagination'][0]
    assert item['Website maximum'] == 3 and item['Maximum verified']
    assert item['Observed maximum'] == 3  # The failed page-4 probe is not valid content.
    assert item['Valid URLs'] == 2 and item['Missing'] == 1 and item['Failed / excluded'] == 1


def test_path_cursor_offset_and_evidence_backed_short_offset():
    rows = [row('/news/page/2/'), row('/api?cursor=abc'), row('/items?offset=20'),
            row('/news-releases?o=10', anchor_text='Next'), row('/unrelated?o=42')]
    analysis = build_analysis(rows, {})
    assert area(analysis, 'Pagination')['Valid URLs'] == 4
    assert len(analysis['pagination']) == 4
    assert all(item['Website maximum'] is None for item in analysis['pagination'])
    assert all(item['Observed maximum'] is None for item in analysis['pagination'] if item['Kind'] != 'Page number')
    assert {item['Kind'] for item in analysis['pagination']} == {'Page number', 'Cursor', 'Offset'}


def test_custom_filter_names_from_recorded_controls_are_scoped():
    controls = [{'source_url': 'https://example.test/catalog', 'form_action': 'https://example.test/items',
                 'control_name': 'sector', 'control_type': 'select', 'form_method': 'get',
                 'request_value': value} for value in ['naval', 'air']]
    analysis = build_analysis([row('/items?sector=naval'), row('/items?sector=air', 'NOT_FOUND'),
                               row('/other?sector=naval')], {'filters': controls})
    assert area(analysis, 'Filters')['Valid URLs'] == 2
    assert area(analysis, 'Filters')['Coverage %'] == 50
    assert {x['Filter state'] for x in analysis['filter_states']} == {'sector=naval', 'sector=air'}
    assert analysis['filter_controls'][0]['Learned values'] == 2
    assert analysis['filter_controls'][0]['Valid URLs'] == 2


def test_browser_url_provenance_does_not_count_dom_or_section_states(tmp_path):
    config = AuditConfig(target_url='https://example.test/', database_path=tmp_path/'audit.sqlite', output_dir=tmp_path/'reports')
    auditor = CoverageAuditor(config)
    seed = config.target_url
    auditor._add_discovery(seed, 'SEED', 0, None, '')
    auditor.db.add_http_result(auditor.run_id, seed, validity='VALID_PAGE', status_code=200)
    states = [{'source_url': seed, 'state_url': url, 'control': 'sector', 'value': 'air', 'label': 'Air', 'method': 'javascript', 'html': '', **extra}
              for url, extra in [(seed+'filtered-air', {}), (seed+'#section', {}), (seed+'tab-page', {'kind': 'content_tab'})]]
    class Browser:
        async def explore(self, url):
            return BrowserDiscovery((Link(seed+'next-results', 'Next', 'pagination'),), (), states=tuple(states))
    class Client:
        def browser_cookies(self, url):
            return []
    async def allowed(url):
        return True
    auditor._allowed = allowed
    try:
        asyncio.run(auditor._crawl_browser(Browser(), Client()))
        rows = {item['url']: item for item in auditor.db.report_rows(auditor.run_id)}
        assert 'PAGINATION' in json.loads(rows[seed+'next-results']['discovery_methods'])
        assert 'FILTER_BROWSER' in json.loads(rows[seed+'filtered-air']['discovery_methods'])
        assert seed+'#section' not in rows and seed+'tab-page' not in rows
    finally:
        auditor.db.close()


def test_dom_only_control_and_loaded_content_tabs_are_not_fabricated_pages():
    controls = [{'source_url': 'https://example.test/', 'control_name': 'year',
                 'control_type': 'rendered', 'form_method': 'javascript', 'request_value': '2021'}]
    analysis = build_analysis([row('/')], {'filters': controls})
    assert area(analysis, 'Filters')['Valid URLs'] == 0
    assert area(analysis, 'Filters')['Coverage %'] is None
    assert next(item for item in analysis['discovery_status'] if item['Area'] == 'Filters')['Status'] == 'Detected; no measurable URLs'
    assert analysis['filter_controls'][0]['Valid URLs'] == 0
    assert 'No independent URL' in analysis['filter_controls'][0]['Assessment']
    assert analysis['filter_states'] == []


def test_empty_scan_areas_explicitly_explain_na():
    analysis = build_analysis([row('/about')], {})
    assert len(analysis['discovery_status']) == 3
    for name in ['Pagination', 'Filters', 'Archives']:
        assert area(analysis, name)['Coverage %'] is None
        status = next(item for item in analysis['discovery_status'] if item['Area'] == name)
        assert status['Status'] == 'No evidence recorded'
        assert 'does not prove' in status['Evidence']
        assert list(area(analysis, name)) == ['Area', 'Valid URLs', 'Found', 'Missing', 'Unverified', 'Coverage %']


def test_archive_date_years_do_not_use_corpus_capture_time():
    assert archive_year('https://example.test/news?publication_year=2021') == '2021'
    assert archive_year('https://example.test/news?date=2022-03-04') == '2022'
    assert archive_year('https://example.test/news?from=2021-01-01&to=2023-01-01') is None
    assert archive_year('https://example.test/news?date=2022-99-99') is None
    rows = [row('/2021/story', 'NOT_FOUND'), row('/2021/other'), row('/2021/pending', None),
            row('/about', evidence_json=json.dumps({'published_at': '2020-01-01'}))]
    analysis = build_analysis(rows, {})
    assert analysis['years'][0]['Coverage %'] == 50
    assert analysis['years'][0]['Unverified'] == 1
    assert area(analysis, 'Archives')['Valid URLs'] == 3
    assert archive_coverage(rows)[0]['coverage_percent'] == 50


def test_excel_has_discovery_assessment_and_enriched_coverage(tmp_path):
    file = tmp_path / 'report.xlsx'
    write_excel(file, {'target': 'https://example.test/'}, [row('/news?page=2&year=2021')],
                [], [], [], [], [], [], [], [])
    workbook = load_workbook(file)
    assert {'Coverage by Area', 'Discovery Assessment', 'Filter Coverage', 'Filter Controls'} <= set(workbook.sheetnames)
    assert 'Coverage %' in [x.value for x in workbook['Pagination Analysis'][1]]
    assert 'Website maximum' in [x.value for x in workbook['Pagination Analysis'][1]]
    assert 'Year' in [x.value for x in workbook['Archive Coverage'][1]]


def test_refresh_reports_uses_saved_data_without_corpus_or_website(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError('Report refresh must not contact website or Corpus')
    monkeypatch.setattr('httpx.AsyncClient.send', forbidden)
    database = tmp_path / 'audit.sqlite'
    db = AuditDatabase(database)
    run_id = db.create_run('https://example.test/', 'https://api.test/', {})
    url = 'https://example.test/news?page=2'
    db.add_url(run_id, DiscoveredURL(url, url, 'STATIC', 1))
    db.add_http_result(run_id, url, status_code=200, validity='VALID_PAGE')
    db.add_corpus_match(run_id, url, 'NOT_FOUND', None, None, None, None)
    db.finish_run(run_id)
    original = database.read_bytes()
    refresh(database, run_id, tmp_path / 'reports')
    assert database.read_bytes() == original
    assert (tmp_path / 'reports' / 'coverage_report.html').is_file()
    db.connection.execute("UPDATE audit_runs SET status='RUNNING' WHERE id=?", (run_id,))
    db.connection.commit()
    with pytest.raises(ValueError, match='finish or pause'):
        refresh(database, run_id, tmp_path / 'active-reports')
    assert not (tmp_path / 'active-reports').exists()
    db.close()


def test_offline_report_discovery_sections_render_empty_and_measured(tmp_path):
    empty = tmp_path / 'empty.html'
    measured = tmp_path / 'measured.html'
    write_html(empty, {'target': 'https://example.test/'}, [row('/about')], [])
    write_html(measured, {'target': 'https://example.test/'}, [row('/news?page=2&year=2021'),
               row('/news?page=3&year=2021', 'NOT_FOUND')], [])
    output = Path(__file__).resolve().parents[1] / 'reports/ui-review'
    output.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.route('http://**/*', lambda route: route.abort())
        page.route('https://**/*', lambda route: route.abort())
        page.goto(empty.as_uri())
        expect(page.locator('#area-table th')).to_have_count(6)
        expect(page.locator('#area-table')).to_contain_text('N/A')
        page.get_by_role('button', name='Discovery analysis').click()
        expect(page.locator('#pagination-table')).to_contain_text('N/A')
        expect(page.locator('#filter-controls')).to_contain_text('Same-page content tabs')
        page.screenshot(path=str(output / 'discovery-empty.png'), full_page=True)
        page.goto(measured.as_uri())
        page.get_by_role('button', name='Discovery analysis').click()
        expect(page.locator('#pagination-table')).to_contain_text('Unknown')
        expect(page.locator('#filter-coverage')).to_contain_text('year=2021')
        expect(page.locator('#archive-table')).to_contain_text('2021')
        for width in [1440, 390]:
            page.set_viewport_size({'width': width, 'height': 1000})
            page.screenshot(path=str(output / f'discovery-measured-{width}.png'), full_page=True)
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
        assert errors == []
        browser.close()
