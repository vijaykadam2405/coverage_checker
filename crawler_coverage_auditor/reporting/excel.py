from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from crawler_coverage_auditor.analysis.coverage import is_inventory_url


SHEETS = ("Executive Summary", "All Discovered URLs", "Missing From Corpus", "Partial Text", "Depth Coverage",
          "Pagination Analysis", "Filter Analysis", "Archive Coverage", "Sitemap Coverage", "Documents",
          "JS Only URLs", "Redirects & Canonicals", "Blocked And Errors", "Structural Gaps", "Evidence")


def _write_table(sheet, rows: list[dict]) -> None:
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        values = [row.get(column) for column in headers]
        sheet.append(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value for value in values])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(58, max(12, max(len(str(cell.value or "")) for cell in column) + 2))


def write_excel(path: Path, summary: dict, rows: list[dict], depth_rows: list[dict], issues: list[dict], pagination_rows: list[dict],
                filter_rows: list[dict], document_rows: list[dict], archive_rows: list[dict], sitemap_rows: list[dict], js_rows: list[dict], *,
                limits: list[dict] | None = None, events: list[dict] | None = None, errors: list[dict] | None = None) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Executive Summary"
    _write_table(first, [summary])
    mappings = {
        "All Discovered URLs": rows,
        "Missing From Corpus": [row for row in rows if row.get("match_status") == "NOT_FOUND" and is_inventory_url(row)],
        "Partial Text": [row for row in rows if row.get('page_type', 'HTML') == 'HTML' and row.get("text_status") in {"FOUND_PARTIAL_TEXT", "FOUND_VERY_PARTIAL_TEXT"}],
        "Depth Coverage": depth_rows,
        "Pagination Analysis": pagination_rows,
        "Filter Analysis": filter_rows, "Archive Coverage": archive_rows, "Sitemap Coverage": sitemap_rows, "Documents": document_rows, "JS Only URLs": js_rows,
        "Redirects & Canonicals": [row for row in rows if row.get("canonical_url") or row.get("final_url")],
        "Blocked And Errors": [row for row in rows if row.get("validity") not in {None, "VALID_PAGE", "URL_ONLY"}],
        "Structural Gaps": issues,
        "Evidence": rows,
    }
    for name in SHEETS[1:]:
        _write_table(workbook.create_sheet(name), mappings[name])
    for name, records in [("Coverage Boundaries", limits or []), ("Activity", events or []), ("Errors", errors or []),
                          ("Unverified Corpus", [row for row in rows if is_inventory_url(row) and row.get("match_status") in {None, "REVIEW_REQUIRED"}])]:
        _write_table(workbook.create_sheet(name), records)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
